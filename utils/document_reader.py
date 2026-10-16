import os
from typing import Dict, List, Optional

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import pypdf
    PdfReader = pypdf.PdfReader
except ImportError:
    try:
        import PyPDF2
        PdfReader = PyPDF2.PdfReader
    except ImportError:
        PdfReader = None


def baca_pdf(filepath: str) -> str:
    """Membaca dan mengekstrak seluruh teks dari file PDF."""
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File PDF tidak ditemukan di: {filepath}")

    if PdfReader is None:
        raise ImportError(
            "Package 'pypdf' atau 'PyPDF2' diperlukan untuk membaca file PDF. "
            "Silakan instal dengan: pip install PyPDF2"
        )

    teks_halaman = []
    with open(filepath, "rb") as f:
        reader = PdfReader(f)
        for idx, page in enumerate(reader.pages, 1):
            ekstraksi = page.extract_text()
            if ekstraksi and ekstraksi.strip():
                teks_halaman.append(ekstraksi.strip())

    return "\n\n".join(teks_halaman)


def baca_rubrik(filepath: str) -> str:
    """Membaca isi file rubrik dari format .docx, .txt, atau .md.

    Args:
        filepath: Path lengkap atau relatif menuju file rubrik.

    Returns:
        Isi teks dari file rubrik.

    Raises:
        FileNotFoundError: Jika file tidak ditemukan.
        ValueError: Jika format ekstensi tidak didukung.
        ImportError: Jika format docx dibuka namun python-docx belum terinstall.
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"File rubrik tidak ditemukan di path: {filepath}")

    ext = os.path.splitext(filepath)[1].lower()

    if ext in [".txt", ".md"]:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            return f.read().strip()
    elif ext == ".docx":
        if Document is None:
            raise ImportError(
                "Package 'python-docx' diperlukan untuk membaca file .docx. "
                "Silakan install dengan: pip install python-docx"
            )
        doc = Document(filepath)
        teks = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(teks).strip()
    else:
        raise ValueError(
            f"Format file rubrik '{ext}' tidak didukung. "
            "Gunakan salah satu dari format berikut: .docx, .txt, atau .md."
        )


def ambil_rubrik_dari_folder(folder_path: str, nama_file: Optional[str] = None) -> str:
    """Mengambil dan membaca teks rubrik dari folder prompts.

    Jika nama_file diberikan, file tersebut yang dibaca.
    Jika nama_file kosong, akan mencari file rubrik pertama yang didukung (.docx, .md, .txt).

    Args:
        folder_path: Path folder prompts.
        nama_file: Nama file opsional (misal: 'rubrik_dengan_referensi.md').

    Returns:
        String isi rubrik yang terbaca.
    """
    if not os.path.exists(folder_path):
        os.makedirs(folder_path, exist_ok=True)
        raise FileNotFoundError(
            f"Folder '{folder_path}' belum memiliki file rubrik. "
            "Silakan tambahkan file rubrik (.docx / .txt / .md) ke dalam folder tersebut."
        )

    if nama_file:
        target = os.path.join(folder_path, nama_file)
        return baca_rubrik(target)

    # Cari file yang didukung dalam folder
    file_didukung = [".docx", ".md", ".txt"]
    for file in sorted(os.listdir(folder_path)):
        ext = os.path.splitext(file)[1].lower()
        if ext in file_didukung:
            return baca_rubrik(os.path.join(folder_path, file))

    raise FileNotFoundError(
        f"Tidak ditemukan file rubrik (.docx, .md, .txt) di dalam folder '{folder_path}'."
    )


def baca_semua_referensi(folder_path: str) -> Dict[str, any]:
    """Membaca dan mengelompokkan dokumen referensi/soal di folder referensi.

    Mendukung dokumen teks: .txt, .md, .docx, .pdf.
    Mendukung gambar acuan: .jpg, .jpeg, .png, .webp, .jpe.

    Args:
        folder_path: Path direktori referensi.

    Returns:
        Dict dengan kunci:
            - 'teks': Gabungan teks dari semua dokumen narasi/soal.
            - 'gambar': List path file gambar acuan.
    """
    if not os.path.exists(folder_path):
        os.makedirs(folder_path, exist_ok=True)
        return {"teks": "", "gambar": []}

    referensi_teks_list: List[str] = []
    referensi_gambar_list: List[str] = []

    ekstensi_teks = [".txt", ".md", ".docx", ".pdf"]
    ekstensi_gambar = [".jpg", ".jpeg", ".png", ".webp", ".jpe"]

    for filename in sorted(os.listdir(folder_path)):
        if filename.startswith("."):
            continue

        filepath = os.path.join(folder_path, filename)
        if not os.path.isfile(filepath):
            continue

        ext = os.path.splitext(filename)[1].lower()

        try:
            if ext in [".txt", ".md"]:
                with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                    konten = f.read().strip()
                if konten:
                    referensi_teks_list.append(f"--- [Dokumen Soal/Acuan: {filename}] ---\n{konten}")

            elif ext == ".docx":
                if Document is not None:
                    doc = Document(filepath)
                    paragraf = [p.text for p in doc.paragraphs if p.text.strip()]
                    konten = "\n".join(paragraf).strip()
                    if konten:
                        referensi_teks_list.append(f"--- [Dokumen Soal/Acuan: {filename}] ---\n{konten}")
                else:
                    print(f" [!] Peringatan: python-docx belum terinstall, melewatkan {filename}")

            elif ext == ".pdf":
                konten = baca_pdf(filepath)
                if konten:
                    referensi_teks_list.append(f"--- [Dokumen Soal/Acuan: {filename}] ---\n{konten}")

            elif ext in ekstensi_gambar:
                referensi_gambar_list.append(filepath)

        except Exception as e:
            print(f" [!] Peringatan: Gagal membaca file referensi '{filename}': {e}")

    gabungan_teks = "\n\n".join(referensi_teks_list)
    return {
        "teks": gabungan_teks,
        "gambar": referensi_gambar_list
    }
